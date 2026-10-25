import csv
import io
import re
from copy import copy
from pathlib import Path

import xlrd
from openpyxl import Workbook, load_workbook
from openpyxl.utils.cell import range_boundaries
from openpyxl.utils import get_column_letter

MAX_FILES = 20
MAX_UPLOAD_BYTES = 100 * 1024 * 1024
ALLOWED_EXTENSIONS = {".xlsx", ".xls", ".csv"}


def _safe_filename(value):
    name = str(value or "").replace("\\", "/").rsplit("/", 1)[-1]
    name = re.sub(r"[^A-Za-z0-9._-]", "_", name)
    if not name or Path(name).suffix.lower() not in ALLOWED_EXTENSIONS:
        raise ValueError(f"Format file tidak didukung: {name or 'tanpa nama'}")
    return name


def _trim_row(row):
    values = list(row)
    while values and values[-1] in (None, ""):
        values.pop()
    return values


def _header_key(value):
    return " ".join(str(value or "").strip().casefold().split())


def _find_header(rows, filename):
    for index, row in enumerate(rows):
        values = _trim_row(row)
        if not any(value not in (None, "") for value in values):
            continue
        keys = [_header_key(value) for value in values]
        if any(not key for key in keys):
            raise ValueError(f"Header kosong ditemukan pada {filename}")
        if len(set(keys)) != len(keys):
            raise ValueError(f"Header duplikat ditemukan pada {filename}")
        return index, values, keys
    raise ValueError(f"Header tidak ditemukan pada {filename}")


def _read_csv(filename, content):
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError:
        try:
            text = content.decode("cp1252")
        except UnicodeDecodeError as error:
            raise ValueError(f"Encoding CSV tidak didukung: {filename}") from error
    sample = text[:8192]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    return [row for row in csv.reader(io.StringIO(text), dialect) if row]


def _read_xlsx(filename, content):
    try:
        workbook = load_workbook(io.BytesIO(content), data_only=False)
    except Exception as error:
        raise ValueError(f"File Excel tidak dapat dibaca: {filename}") from error
    if not workbook.worksheets:
        workbook.close()
        raise ValueError(f"Workbook tidak memiliki sheet: {filename}")
    worksheet = next(
        (sheet for sheet in workbook.worksheets if any(cell.value is not None for row in sheet.iter_rows() for cell in row)),
        workbook.worksheets[0],
    )
    rows = [_trim_row(row) for row in worksheet.iter_rows(values_only=True)]
    return rows, workbook, worksheet


def _read_xls(filename, content):
    try:
        workbook = xlrd.open_workbook(file_contents=content, on_demand=True)
        if not workbook.nsheets:
            raise ValueError(f"Workbook tidak memiliki sheet: {filename}")
        sheet = workbook.sheet_by_index(0)
        rows = []
        for row_index in range(sheet.nrows):
            row = []
            for column_index in range(sheet.ncols):
                cell = sheet.cell(row_index, column_index)
                if cell.ctype == xlrd.XL_CELL_EMPTY:
                    row.append(None)
                elif cell.ctype == xlrd.XL_CELL_DATE:
                    row.append(xlrd.xldate_as_datetime(cell.value, workbook.datemode))
                else:
                    row.append(cell.value)
            rows.append(_trim_row(row))
        workbook.release_resources()
        return rows
    except Exception as error:
        if isinstance(error, ValueError):
            raise
        raise ValueError(f"File Excel tidak dapat dibaca: {filename}") from error


def _read_source(filename, content):
    extension = Path(filename).suffix.lower()
    if extension == ".csv":
        rows = _read_csv(filename, content)
        workbook = worksheet = None
    elif extension == ".xlsx":
        rows, workbook, worksheet = _read_xlsx(filename, content)
    else:
        rows = _read_xls(filename, content)
        workbook = worksheet = None

    try:
        header_index, headers, keys = _find_header(rows, filename)
        data_rows = []
        for row in rows[header_index + 1:]:
            values = _trim_row(row)
            if any(value not in (None, "") for value in values):
                if len(values) > len(headers) and any(value not in (None, "") for value in values[len(headers):]):
                    raise ValueError(f"Data memiliki kolom tanpa header pada {filename}")
                data_rows.append(values[:len(headers)])
        return {
            "filename": filename,
            "extension": extension,
            "headers": headers,
            "keys": keys,
            "rows": data_rows,
            "workbook": workbook,
            "worksheet": worksheet,
            "header_index": header_index,
        }
    except Exception:
        if workbook:
            workbook.close()
        raise


def compile_workbooks(files):
    if not 2 <= len(files) <= MAX_FILES:
        raise ValueError(f"Pilih antara 2 sampai {MAX_FILES} file")

    prepared = []
    total_size = 0
    for filename, content in files:
        safe_name = _safe_filename(filename)
        if not isinstance(content, bytes) or not content:
            raise ValueError(f"File kosong atau tidak valid: {safe_name}")
        total_size += len(content)
        if total_size > MAX_UPLOAD_BYTES:
            raise ValueError("Total ukuran file melebihi batas 100 MB")
        prepared.append((safe_name, content))

    records = []
    try:
        for filename, content in prepared:
            records.append(_read_source(filename, content))

        template_index = next(
            (index for index, record in enumerate(records) if Path(record["filename"]).suffix.lower() == ".xlsx"),
            None,
        )
        if template_index is None:
            raise ValueError("Sertakan minimal satu file .xlsx sebagai template agar format tetap terjaga")
        template = records[template_index]
        expected_keys = template["keys"]
        for record in records:
            if set(record["keys"]) != set(expected_keys):
                missing = [key for key in expected_keys if key not in record["keys"]]
                extra = [key for key in record["keys"] if key not in expected_keys]
                detail = []
                if missing:
                    detail.append("kolom hilang: " + ", ".join(missing))
                if extra:
                    detail.append("kolom tambahan: " + ", ".join(extra))
                raise ValueError(f"Header {record['filename']} tidak cocok ({'; '.join(detail)})")

        template_workbook = template["workbook"]
        template_sheet = template["worksheet"]
        if template_workbook:
            header_row = template["header_index"] + 1
            style_row = header_row + 1 if header_row < template_sheet.max_row else None
            styles = []
            if style_row:
                styles = [copy(template_sheet.cell(style_row, column)._style) for column in range(1, len(template["headers"]) + 1)]
                row_height = template_sheet.row_dimensions[style_row].height
            else:
                row_height = None
            if template_sheet.max_row > header_row:
                template_sheet.delete_rows(header_row + 1, template_sheet.max_row - header_row)
            workbook = template_workbook
            worksheet = template_sheet
            for sheet in list(workbook.worksheets):
                if sheet is not worksheet:
                    workbook.remove(sheet)
        else:
            workbook = Workbook()
            worksheet = workbook.active
            worksheet.title = Path(template["filename"]).stem[:31] or "Compiled"
            header_row = 1
            styles = []
            row_height = None
            worksheet.append(template["headers"])

        for record in records:
            source_columns = {key: index for index, key in enumerate(record["keys"])}
            for row in record["rows"]:
                output_row = [row[source_columns[key]] if source_columns[key] < len(row) else None for key in expected_keys]
                worksheet.append(output_row)
                output_index = worksheet.max_row
                if record["extension"] == ".csv":
                    for column, value in enumerate(output_row, start=1):
                        if isinstance(value, str) and value.startswith("="):
                            worksheet.cell(output_index, column).data_type = "s"
                if styles:
                    for column, style in enumerate(styles, start=1):
                        if style is not None:
                            worksheet.cell(output_index, column)._style = copy(style)
                    if row_height is not None:
                        worksheet.row_dimensions[output_index].height = row_height

        if worksheet.auto_filter.ref:
            min_column, min_row, max_column, _ = range_boundaries(worksheet.auto_filter.ref)
            worksheet.auto_filter.ref = (
                f"{get_column_letter(min_column)}{min_row}:"
                f"{get_column_letter(max_column)}{worksheet.max_row}"
            )
        for table in worksheet.tables.values():
            min_column, min_row, max_column, _ = range_boundaries(table.ref)
            if min_row == header_row:
                table.ref = (
                    f"{get_column_letter(min_column)}{min_row}:"
                    f"{get_column_letter(max_column)}{worksheet.max_row}"
                )
                if table.autoFilter:
                    table.autoFilter.ref = table.ref

        output = io.BytesIO()
        workbook.save(output)
        return output.getvalue()
    finally:
        closed = set()
        for record in records:
            workbook = record["workbook"]
            if workbook and id(workbook) not in closed:
                workbook.close()
                closed.add(id(workbook))
