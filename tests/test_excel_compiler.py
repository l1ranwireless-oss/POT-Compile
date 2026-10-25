import io
import unittest

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

from excel_compiler import compile_workbooks


def make_xlsx(rows, styled=False):
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Source"
    for row in rows:
        worksheet.append(row)
    if styled:
        worksheet["A1"].font = Font(bold=True, color="FFFFFF")
        worksheet["A1"].fill = PatternFill("solid", fgColor="1E40AF")
        worksheet["B2"].number_format = "0.00"
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


class ExcelCompilerTests(unittest.TestCase):
    def test_compiles_csv_into_xlsx_template_and_preserves_style(self):
        template = make_xlsx([["Name", "Amount"], ["Excel row", 1]], styled=True)
        csv_content = b"amount,name\n2,CSV row\n"

        output = compile_workbooks([
            ("input.csv", csv_content),
            ("template.xlsx", template),
        ])

        workbook = load_workbook(io.BytesIO(output), data_only=False)
        self.assertEqual(workbook.sheetnames, ["Source"])
        worksheet = workbook.active
        self.assertEqual(list(worksheet.values), [
            ("Name", "Amount"),
            ("CSV row", "2"),
            ("Excel row", 1),
        ])
        self.assertTrue(worksheet["A1"].font.bold)
        self.assertEqual(worksheet["A1"].fill.fgColor.rgb, "001E40AF")
        self.assertEqual(worksheet["B2"].number_format, "0.00")
        self.assertEqual(worksheet["B3"].number_format, "0.00")
        workbook.close()

    def test_rejects_incompatible_headers(self):
        first = make_xlsx([["Name", "Amount"], ["Alpha", 1]])
        second = b"name,quantity\nBeta,2\n"

        with self.assertRaisesRegex(ValueError, "Header input.csv tidak cocok"):
            compile_workbooks([("first.xlsx", first), ("input.csv", second)])

    def test_requires_multiple_files(self):
        with self.assertRaisesRegex(ValueError, "antara 2 sampai"):
            compile_workbooks([("input.csv", b"name\nAlpha\n")])

    def test_requires_xlsx_template_to_preserve_format(self):
        with self.assertRaisesRegex(ValueError, "minimal satu file .xlsx"):
            compile_workbooks([
                ("first.csv", b"name\nAlpha\n"),
                ("second.csv", b"name\nBeta\n"),
            ])

    def test_csv_formula_text_is_not_written_as_excel_formula(self):
        template = make_xlsx([["Name"], ["Template"]])
        output = compile_workbooks([
            ("template.xlsx", template),
            ("input.csv", b"Name\n=1+1\n"),
        ])

        workbook = load_workbook(io.BytesIO(output), data_only=False)
        cell = workbook.active["A3"]
        self.assertEqual(cell.value, "=1+1")
        self.assertEqual(cell.data_type, "s")
        workbook.close()


if __name__ == "__main__":
    unittest.main()
