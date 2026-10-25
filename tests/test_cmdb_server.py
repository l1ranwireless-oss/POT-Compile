import io
import json
import tempfile
import threading
import unittest
from http.server import ThreadingHTTPServer
from urllib.error import HTTPError
from urllib.request import Request, urlopen

from openpyxl import Workbook, load_workbook

from cmdb_server import make_handler


def workbook_bytes(rows):
    workbook = Workbook()
    worksheet = workbook.active
    for row in rows:
        worksheet.append(row)
    output = io.BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


class CMDBServerTests(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.token = "test-token-" + ("x" * 32)
        self.server = ThreadingHTTPServer(
            ("127.0.0.1", 0), make_handler(self.directory.name, self.token)
        )
        self.base_url = f"http://127.0.0.1:{self.server.server_port}"
        self.thread = threading.Thread(target=self.server.serve_forever, daemon=True)
        self.thread.start()

    def tearDown(self):
        self.server.shutdown()
        self.server.server_close()
        self.thread.join()
        self.directory.cleanup()

    def request(self, path, method="GET", body=None, filename=None, authorized=True):
        headers = {}
        if authorized:
            headers["Authorization"] = f"Bearer {self.token}"
        if filename:
            headers["X-CMDB-Filename"] = filename
        request = Request(
            self.base_url + path,
            data=body,
            headers=headers,
            method=method,
        )
        return urlopen(request)

    def test_requires_authentication(self):
        with self.assertRaises(HTTPError) as raised:
            self.request("/api/cmdb", authorized=False)
        self.assertEqual(raised.exception.code, 401)
        raised.exception.close()

    def test_excel_compile_requires_authentication(self):
        request = Request(
            self.base_url + "/api/excel/compile",
            data=b"not multipart",
            headers={"Content-Type": "application/octet-stream"},
            method="POST",
        )
        with self.assertRaises(HTTPError) as raised:
            urlopen(request)
        self.assertEqual(raised.exception.code, 401)
        raised.exception.close()

    def test_excel_compile_returns_downloadable_workbook(self):
        boundary = "excel-compiler-test-boundary"
        files = [
            ("template.xlsx", workbook_bytes([["Name", "Amount"], ["Alpha", 1]])),
            ("extra.csv", b"amount,name\n2,Beta\n"),
        ]
        body = b"".join(
            b"--" + boundary.encode() + b"\r\n"
            + f'Content-Disposition: form-data; name="files"; filename="{name}"\r\n'.encode()
            + b"Content-Type: application/octet-stream\r\n\r\n"
            + content + b"\r\n"
            for name, content in files
        ) + b"--" + boundary.encode() + b"--\r\n"
        request = Request(
            self.base_url + "/api/excel/compile",
            data=body,
            headers={
                "Authorization": f"Bearer {self.token}",
                "Content-Type": f"multipart/form-data; boundary={boundary}",
            },
            method="POST",
        )

        response = urlopen(request)
        self.assertEqual(response.headers.get_content_type(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        workbook = load_workbook(io.BytesIO(response.read()), read_only=True)
        self.assertEqual(workbook.sheetnames, ["Sheet"])
        self.assertEqual(list(workbook.active.values), [
            ("Name", "Amount"),
            ("Alpha", 1),
            ("Beta", "2"),
        ])
        workbook.close()

    def test_rejects_malformed_metadata(self):
        with open(f"{self.directory.name}/metadata.json", "w", encoding="utf-8") as metadata:
            json.dump([], metadata)

        with self.assertRaises(HTTPError) as raised:
            self.request("/api/cmdb")
        self.assertEqual(raised.exception.code, 500)
        self.assertEqual(json.loads(raised.exception.read())["error"], "Metadata CMDB server tidak valid")
        raised.exception.close()

    def test_save_update_and_load_source(self):
        content = b"CMDB source bytes"
        response = self.request(
            "/api/cmdb/save",
            method="POST",
            body=content,
            filename="cmdb.xlsx",
        )
        self.assertEqual(json.loads(response.read())["size"], len(content))

        with self.assertRaises(HTTPError) as raised:
            self.request(
                "/api/cmdb/save",
                method="POST",
                body=content,
                filename="cmdb.xlsx",
            )
        self.assertEqual(raised.exception.code, 409)
        raised.exception.close()

        replacement = b"Updated CMDB source"
        self.request(
            "/api/cmdb/update",
            method="PUT",
            body=replacement,
            filename="cmdb-latest.xlsx",
        ).close()

        status = json.loads(self.request("/api/cmdb").read())
        self.assertEqual(status["name"], "cmdb-latest.xlsx")
        self.assertEqual(status["size"], len(replacement))
        self.assertEqual(self.request("/api/cmdb/file").read(), replacement)

    def test_rejects_unsupported_extension(self):
        with self.assertRaises(HTTPError) as raised:
            self.request(
                "/api/cmdb/save",
                method="POST",
                body=b"not a CMDB",
                filename="payload.exe",
            )
        self.assertEqual(raised.exception.code, 400)
        raised.exception.close()

    def test_saved_source_survives_server_restart(self):
        content = b"Persistent CMDB"
        self.request(
            "/api/cmdb/save",
            method="POST",
            body=content,
            filename="cmdb.csv",
        ).close()
        self.server.shutdown()
        self.server.server_close()
        self.thread.join()

        self.server = ThreadingHTTPServer(
            ("127.0.0.1", 0), make_handler(self.directory.name, self.token)
        )
        self.base_url = f"http://127.0.0.1:{self.server.server_port}"
        self.thread = threading.Thread(target=self.server.serve_forever, daemon=True)
        self.thread.start()

        status = json.loads(self.request("/api/cmdb").read())
        self.assertTrue(status["exists"])
        self.assertEqual(self.request("/api/cmdb/file").read(), content)


if __name__ == "__main__":
    unittest.main()
